"""Build the preconstruction Dashboard: a workbook of Excel tables + live formulas, and the nightly data file.

    python3 dashboard/build_dashboard.py --bid-log "Bid Log 2026.csv" --bid-metrics "Bid Metrics 2026.csv" \
        [--prices "Price Movers.csv"] [--job-cost "Job Cost History.csv"] \
        [--out Dashboard.xlsx] [--data "Dashboard Data.txt"] [--previous DIR] [--no-recalc]

Two outputs:
* Dashboard.xlsx — the workbook (tables tblBidLog, tblGCs, tblJobTypes, tblLabor, tblJobCost, tblMovers; every summary
  is a formula over those tables). Uploaded to SharePoint ONCE by hand; after that it is never replaced.
* Dashboard Data 1.txt … Dashboard Data 4.txt — the rows for each table as JSON, split into PARTS parts of at most
  PART_MAX bytes each (System Rules C2: a run never writes a file over 20 KB). A table too big for one part (tblBidLog)
  continues in the next part. Every part carries the same "set" hash, so the Office Script refuses a mix of nights;
  the stamp and About text ride in the last part. Unused parts are written empty (still valid). Runs upload the last part nightly
  and the other parts only when they changed (--previous DIR = the current parts copied from SharePoint).
  A Power Automate flow (Flow 3) passes the four files to dashboard/load_dashboard_data.ts (an Office Script).
Formula columns inside tables (in_period, the GC and job-type stats) are kept by the script, never sent as data.

Checks: LibreOffice recalculates a copy, every formula is scanned for errors, and hit rate and labor $/SF are
re-computed in Python and compared. Last line says whether to upload. Exit 1 = do not upload.
Hit rate = won / (won + lost) by bid date; submitted-awaiting and closed are shown, never in the rate.
"""
import argparse, csv, datetime as dt, json, os, re, shutil, statistics, subprocess, sys, tempfile
from collections import Counter, defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo

UNIT_TYPES = [("cmu_gray", "Gray CMU"), ("cmu_arch", "Architectural CMU"), ("brick", "Brick"),
              ("stone", "Stone"), ("cast_stone", "Cast stone")]
PER = [("sf", "#,##0"), ("labor_per_sf", "$#,##0.00"), ("sf_per_manday", "#,##0.0"),
       ("material_per_sf", "$#,##0.00"), ("sell_per_sf", "$#,##0.00")]
STD_WAGES = "Bricklayer $30 · Foreman $65 · Laborer $20 · Operator $27 /hr"
PERIODS = ["All", "YTD", "Trailing 12 months"]
GC_SUFFIX = {"llc", "inc", "lp", "i", "ltd", "co", "corp", "company", "construction", "constructors", "contractors",
             "contracting", "general", "group", "builders", "national", "services", "the"}
GC_KEEP_APART = {"arco murray", "rudick joeris", "evans consulting"}   # different firms the suffix rule would merge
N_JOBTYPES, N_TOPGC, N_MOVERS = 10, 12, 15

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True)
F_SECTION = Font(name=FONT, size=12, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="666666")
F_INPUT = Font(name=FONT, size=10, bold=True, color="0000FF")
FILL_HEAD = PatternFill("solid", fgColor="1F2937")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(bottom=Side(style="thin", color="D1D5DB"))
MONEY, RATE, PCT, INT, DATE = '$#,##0;($#,##0);"-"', '$#,##0.00;($#,##0.00);"-"', '0.0%;(0.0%);"-"', '#,##0;(#,##0);"-"', "yyyy-mm-dd"
EPOCH = dt.date(1899, 12, 30)


# ---------- input ----------
def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def num(v):
    try:
        return float(v) if str(v).strip() != "" else None
    except ValueError:
        return None


def day(v):
    try:
        return dt.date.fromisoformat(str(v).strip()[:10])
    except ValueError:
        return None


def serial(d):
    return (d - EPOCH).days if d else None


def gc_key(name):
    s = re.sub(r"\(.*?\)", " ", name or "").lower().replace("&", " and ")
    toks = [t for t in re.split(r"[^a-z0-9]+", s) if t]
    core = " ".join(t for t in toks if t not in GC_SUFFIX) or " ".join(toks)
    for k in GC_KEEP_APART:
        if core.startswith(k):
            return k
    return core


def gc_groups(rows):
    """raw GC name -> display group (most common spelling, parentheses dropped; ties: shortest)."""
    by = defaultdict(Counter)
    for r in rows:
        g = (r["gc"] or "").strip()
        if g and g.lower() != "unknown":
            by[gc_key(g)][g] += 1
    raw2grp, merges = {}, []
    for k, c in by.items():
        clean = Counter()
        for raw, n in c.items():
            clean[re.sub(r"\s*\(.*?\)\s*", " ", raw).strip()] += n
        disp = sorted(clean.items(), key=lambda t: (-t[1], len(t[0])))[0][0]
        for raw in c:
            raw2grp[raw] = disp
        if len(c) > 1 or disp != next(iter(c)):
            merges.append((disp, sorted(c)))
    return raw2grp, sorted(merges)


# ---------- data (shared by the workbook and the data file) ----------
def prepare(a):
    bl = read_csv(a.bid_log)
    bm = read_csv(a.bid_metrics)
    movers = [p for p in read_csv(a.prices) if p.get("mover") == "y"] if a.prices and os.path.exists(a.prices) else []
    jc = read_csv(a.job_cost) if a.job_cost and os.path.exists(a.job_cost) else []
    warn = []
    if not (a.prices and os.path.exists(a.prices)):
        warn.append("no Price Movers.csv given — movers section empty")
    raw2grp, merges = gc_groups(bl)
    T = {}

    # tblBidLog — one row per bid per GC, plus project-level columns so a project bid to many GCs counts once:
    # project_first = 1 on the project's earliest row; project_status = awarded if any GC awarded it to us,
    # open while any GC is still out, lost if decided against us, closed if no result; project_amount = its price.
    norm = lambda x: re.sub(r"[^a-z0-9]+", " ", (x or "").lower()).strip()
    groups = defaultdict(list)
    for i, r in enumerate(bl):
        groups[(norm(r["project"]), norm(r["city"]))].append(i)
    pfirst, pstatus, pamount = {}, {}, {}
    for idx in groups.values():
        st = {bl[i]["status"] for i in idx}
        first = min(idx, key=lambda i: (bl[i]["bid_date"] or "9999", i))
        status = ("awarded" if "awarded" in st else "open" if "submitted" in st or ("bidding" in st and len(st) > 1)
                  else "lost" if "lost" in st else "closed" if "closed" in st else sorted(st)[0])
        amts = [num(bl[i]["submitted_amount"]) for i in idx if num(bl[i]["submitted_amount"])]
        for i in idx:
            pfirst[i], pstatus[i] = int(i == first), status
        pamount[first] = max(amts) if amts else None
    P_MULTI = sum(1 for idx in groups.values() if len(idx) > 1)
    cols = ["bid_id", "project", "job_type", "gc_group", "bid_date", "status", "submitted_amount",
            "project_first", "project_status", "project_amount"]
    rows = []
    for i, r in enumerate(bl):
        gc = (r["gc"] or "").strip()
        rows.append([r["bid_id"], r["project"], r["job_type"] or "unknown", raw2grp.get(gc, "(unknown)"),
                     serial(day(r["bid_date"])), r["status"], num(r["submitted_amount"]),
                     pfirst[i], pstatus[i], pamount.get(i)])
    T["tblBidLog"] = (cols, rows)

    # tblGCs / tblJobTypes: key column only (stats are table formulas); ordered by all-time bid count
    counted = ("submitted", "awarded", "lost", "closed")
    gcount = Counter(raw2grp.get((r["gc"] or "").strip(), "(unknown)") for r in bl if r["status"] in counted)
    grps = sorted(gcount, key=lambda g: (g == "(unknown)", -gcount[g], g))
    T["tblGCs"] = (["GC"], [[g] for g in grps])
    jcount = Counter(r["job_type"] or "unknown" for r in bl if r["status"] in counted)
    T["tblJobTypes"] = (["Job type"], [[j] for j in sorted(jcount, key=lambda j: (-jcount[j], j))])

    # tblLabor — linked to the Bid Log by bid_id, else bid price within 0.05%, else an unambiguous name match
    priced = [(num(r["submitted_amount"]) or num(r["base_bid"]), r) for r in bl]
    priced = [(p, r) for p, r in priced if p]
    by_id = {r["bid_id"]: r for r in bl if r["bid_id"]}
    toks_of = lambda s: set(re.findall(r"[a-z0-9]+", s.lower()))
    lcols = ["project", "job_number", "job_type", "bid_date", "bid_price", "linked", "wages", "wages_vs_std", "notes"]
    lcols += [f"{u}_{p}" for u, _ in UNIT_TYPES for p, _ in PER]
    lrows = []
    for r in bm:
        link, how = by_id.get(r.get("bid_id", "")), "id"
        if not link and num(r["bid_price"]):
            price = num(r["bid_price"])
            hits = [x for p, x in priced if abs(p - price) <= max(5, price * 0.0005)]
            hits.sort(key=lambda x: -len(toks_of(r["project"]) & toks_of(x["project"])))
            link, how = (hits[0], "price") if hits else (None, "")
        if not link:
            toks = [t for t in re.findall(r"[a-z0-9]+", r["project"].lower()) if len(t) >= 4]
            cands = [x for x in bl if toks and all(t in x["project"].lower() for t in toks)]
            if cands and len({x["project"] for x in cands}) == 1:
                link, how = cands[0], "name"
        if not link:
            warn.append(f"Labor History: {r['project']} not linked to a Bid Log row"
                        + (" (iteration)" if "iteration" in r.get("notes", "") else ""))
        row = [r["project"], r["job_number"], (link or {}).get("job_type") or None, serial(day((link or {}).get("bid_date", ""))),
               num(r["bid_price"]), how or "no", r["wages"], num(r["wages_vs_standard"]), r["notes"] or None]
        for u, _ in UNIT_TYPES:
            sf = num(r.get(f"{u}_sf"))
            row += [num(r.get(f"{u}_{p}")) if sf else None for p, _ in PER]
        lrows.append(row)
    T["tblLabor"] = (lcols, lrows)

    # tblJobCost
    jcols = ["job_number", "estimate", "scope", "through_date", "status", "est_sf", "est_mandays", "act_mandays",
             "mandays_var_pct", "est_labor_base", "act_labor_base", "labor_var_pct", "est_labor_per_sf", "act_labor_per_sf",
             "est_sf_per_manday", "act_sf_per_manday", "est_material", "act_material", "material_var_pct", "labor_basis",
             "hours_by_class", "notes"]
    text = {"job_number", "estimate", "scope", "status", "labor_basis", "hours_by_class", "notes"}
    T["tblJobCost"] = (jcols, [[serial(day(r.get(k))) if k == "through_date" else (r.get(k) or None) if k in text else num(r.get(k))
                                for k in jcols] for r in jc])

    # tblMovers
    mcols = ["item", "change_pct", "unit_price", "prior_price", "unit", "product", "supplier", "quote_date", "prior_date", "project"]
    movers.sort(key=lambda p: -abs(num(p["change_pct"]) or 0))
    T["tblMovers"] = (mcols, [[p["item"], num(p["change_pct"]), num(p["unit_price"]), num(p["prior_price"]), p["unit"],
                                p["product"], p["supplier"], serial(day(p["quote_date"])), serial(day(p["prior_date"])),
                                p["project"]] for p in movers])

    about = [f"Rebuilt {a.stamp} from Bid Log 2026.csv ({len(bl)} rows = {len(groups)} projects, {P_MULTI} bid to more than one GC), Bid Metrics 2026.csv ({len(bm)} estimates), "
             f"Job Cost History.csv ({len(jc)} rows), Price Movers.csv ({len(movers)} movers).",
             "Hit rate (Dashboard, By Job Type) counts each PROJECT once: won if any GC awarded it to us, lost if decided against us, "
             "open while any GC is still out; won ÷ (won + lost), by the project's first bid date. By GC counts each bid to that GC.",
             f"Labor $/SF = Crew Page base wage spread across unit types by line-item labor (before burden and OH&P). Standard wages: {STD_WAGES}.",
             "", "GC names merged (display ← spellings in the Bid Log):"]
    about += [f"{d} ← " + " | ".join(raws) for d, raws in merges] or ["none"]
    about += ["", "⚠ items:"] + (warn or ["none"])
    P_proj = {k: pstatus[idx[0]] for k, idx in groups.items()}
    return dict(T=T, bl=bl, bm=bm, warn=warn, about=about, proj=P_proj, stamp=f"Rebuilt {a.stamp}. Don't edit data tabs — they reload nightly.")


# ---------- workbook ----------
def head(ws, row, col, labels):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def put(ws, ref, value, fmt=None, font=F_BASE):
    ws[ref] = value
    ws[ref].font = font
    if fmt:
        ws[ref].number_format = fmt


def arr(ws, ref, formula, fmt=None):
    ws[ref] = ArrayFormula(ref, formula)
    ws[ref].font = F_BASE
    if fmt:
        ws[ref].number_format = fmt


def this(tbl, col):
    return f"{tbl}[[#This Row],[{col}]]"


def add_table(ws, name, cols, rows, formulas=None, fmts=None, widths=None, start=1):
    """Header + rows (≥ 1 row so the table never empties) + optional formula columns appended on the right."""
    formulas = formulas or []
    headers = cols + [h for h, _ in formulas]
    head(ws, start, 1, headers)
    ws.row_dimensions[start].height = 30
    body = rows or [[None] * len(cols)]
    for i, r in enumerate(body, start=start + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(i, j, v); c.font = F_BASE
            if fmts and cols[j - 1] in fmts:
                c.number_format = fmts[cols[j - 1]]
        for k, (h, f) in enumerate(formulas):
            c = ws.cell(i, len(cols) + 1 + k, f[0]); c.font = F_BASE
            if f[1]:
                c.number_format = f[1]
    ref = f"A{start}:{L(len(headers))}{start + len(body)}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    for j, w in enumerate(widths or [], start=1):
        ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = ws.cell(start + 1, 2)


def stat_formulas(tbl, key, bidcol, project=False):
    """Bids/Won/Lost/... formulas for a key table against tblBidLog. project=True counts each project once."""
    k = this(tbl, key)
    if project:
        cnt = lambda s: (f'COUNTIFS(tblBidLog[project_first],1,tblBidLog[project_status],"{PSTAT[s]}",'
                         f'tblBidLog[in_period],1,tblBidLog[{bidcol}],{k})')
        sm = lambda s: (f'SUMIFS(tblBidLog[project_amount],tblBidLog[project_first],1,tblBidLog[project_status],"{PSTAT[s]}",'
                        f'tblBidLog[in_period],1,tblBidLog[{bidcol}],{k})')
    else:
        cnt = lambda s: f'COUNTIFS(tblBidLog[status],"{s}",tblBidLog[in_period],1,tblBidLog[{bidcol}],{k})'
        sm = lambda s: f'SUMIFS(tblBidLog[submitted_amount],tblBidLog[status],"{s}",tblBidLog[in_period],1,tblBidLog[{bidcol}],{k})'
    w, l, ws_, ls = this(tbl, "Won"), this(tbl, "Lost"), this(tbl, "Won $"), this(tbl, "Lost $")
    return [("Bids", (f"={cnt('submitted')}+{cnt('awarded')}+{cnt('lost')}+{cnt('closed')}", INT)),
            ("Won", (f"={cnt('awarded')}", INT)), ("Lost", (f"={cnt('lost')}", INT)),
            ("Hit rate", (f'=IF({w}+{l}=0,"",{w}/({w}+{l}))', PCT)),
            ("Won $", (f"={sm('awarded')}", MONEY)), ("Lost $", (f"={sm('lost')}", MONEY)),
            ("Hit rate $", (f'=IF({ws_}+{ls}=0,"",{ws_}/({ws_}+{ls}))', PCT)),
            ("Awaiting", (f"={cnt('submitted')}", INT)), ("Closed", (f"={cnt('closed')}", INT)),
            ("Last bid", (f'=IF(COUNTIFS(tblBidLog[{bidcol}],{k},tblBidLog[bid_date],"<>")=0,"",'
                          f'_xlfn.MAXIFS(tblBidLog[bid_date],tblBidLog[{bidcol}],{k}))', DATE))]


PSTAT = {"submitted": "open", "awarded": "awarded", "lost": "lost", "closed": "closed"}
STAT_HEAD = ["Bids", "Won", "Lost", "Hit rate", "Won $", "Lost $", "Hit rate $", "Awaiting", "Closed"]
STAT_FMT = [INT, INT, INT, PCT, MONEY, MONEY, PCT, INT, INT]


def build(a, P):
    T = P["T"]
    wb = Workbook()
    dash = wb.active; dash.title = "Dashboard"
    sh = {n: wb.create_sheet(n) for n in ["By GC", "By Job Type", "Labor History", "Job Cost", "Movers", "Bid Log", "About"]}
    for nm, ref in (("Period", "Dashboard!$C$4"), ("PeriodFrom", "Dashboard!$C$5")):
        dn = DefinedName(nm, attr_text=ref)
        if hasattr(wb.defined_names, "add"):
            wb.defined_names.add(dn)
        else:
            wb.defined_names[nm] = dn

    bd = this("tblBidLog", "bid_date")
    add_table(sh["Bid Log"], "tblBidLog", *T["tblBidLog"],
              formulas=[("in_period", (f'=IF({bd}="",IF(Period="All",1,0),IF({bd}>=PeriodFrom,1,0))', None))],
              fmts={"bid_date": DATE, "submitted_amount": MONEY, "project_amount": MONEY},
              widths=[26, 40, 12, 26, 11, 11, 13, 9, 11, 13, 9])
    add_table(sh["By GC"], "tblGCs", *T["tblGCs"], formulas=stat_formulas("tblGCs", "GC", "gc_group"),
              widths=[34, 8, 8, 8, 10, 13, 13, 10, 9, 8, 12])
    add_table(sh["By Job Type"], "tblJobTypes", *T["tblJobTypes"], formulas=stat_formulas("tblJobTypes", "Job type", "job_type", project=True),
              widths=[18, 8, 8, 8, 10, 13, 13, 10, 9, 8, 12])
    lf = {"bid_date": DATE, "bid_price": MONEY, "wages_vs_std": PCT}
    lf.update({f"{u}_{p}": f for u, _ in UNIT_TYPES for p, f in PER})
    add_table(sh["Labor History"], "tblLabor", *T["tblLabor"], fmts=lf, widths=[24, 8, 12, 11, 12, 8, 34, 9, 24] + [10] * 25)
    jf = {"through_date": DATE}
    jf.update({k: (PCT if k.endswith("var_pct") else RATE if k.endswith("per_sf") else "#,##0.0" if "manday" in k else MONEY)
               for k in T["tblJobCost"][0] if k.startswith(("est_", "act_", "mandays", "labor_var", "material_var"))})
    add_table(sh["Job Cost"], "tblJobCost", *T["tblJobCost"], fmts=jf, widths=[10, 22, 10, 11, 8] + [11] * 15 + [30, 30])
    add_table(sh["Movers"], "tblMovers", *T["tblMovers"],
              fmts={"change_pct": PCT, "unit_price": RATE, "prior_price": RATE, "quote_date": DATE, "prior_date": DATE},
              widths=[18, 9, 11, 11, 6, 40, 28, 11, 11, 28])

    ab = sh["About"]; ab.column_dimensions["A"].width = 120
    for i, t in enumerate(P["about"], start=1):
        put(ab, f"A{i}", t)

    # Dashboard
    ws = dash
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJK", [24, 13, 13, 13, 30, 13, 30, 11, 12, 13, 12]):
        ws.column_dimensions[col].width = w
    put(ws, "A1", "Preconstruction Dashboard", font=F_TITLE)
    put(ws, "A2", P["stamp"], font=F_NOTE)
    put(ws, "A4", "Period (pick)", font=F_BOLD)
    c = ws["C4"]; c.value, c.font, c.fill = a.period, F_INPUT, FILL_INPUT
    dv = DataValidation(type="list", formula1='"' + ",".join(PERIODS) + '"', allow_blank=False)
    ws.add_data_validation(dv); dv.add("C4")
    put(ws, "A5", "From"); put(ws, "C5", '=IF(C4="YTD",DATE(YEAR(TODAY()),1,1),IF(C4="Trailing 12 months",EDATE(TODAY(),-12),DATE(2000,1,1)))', DATE)
    put(ws, "E4", "Period drives hit rates (by bid date). Labor $/SF uses every estimate.", font=F_NOTE)

    r0 = 7
    put(ws, f"A{r0}", "Labor $/SF by unit type — Crew Page base wage, before burden and OH&P", font=F_SECTION)
    head(ws, r0 + 1, 1, ["Unit type", "Median $/SF", "Low $/SF", "Low job", "", "High $/SF", "High job", "Estimates", "Total SF", "Median SF/man-day"])
    ws.merge_cells(start_row=r0 + 1, start_column=4, end_row=r0 + 1, end_column=5)
    for k, (u, name) in enumerate(UNIT_TYPES):
        r = r0 + 2 + k
        lp, sf, spm = f"tblLabor[{u}_labor_per_sf]", f"tblLabor[{u}_sf]", f"tblLabor[{u}_sf_per_manday]"
        job = lambda ref: (f'=IF({ref}="","",INDEX(tblLabor[project],MATCH({ref},{lp},0))'
                           f'&IF(INDEX(tblLabor[notes],MATCH({ref},{lp},0))="",""," ("&INDEX(tblLabor[notes],MATCH({ref},{lp},0))&")"))')
        put(ws, f"A{r}", name, font=F_BOLD)
        arr(ws, f"B{r}", f'=IF(COUNT({lp})=0,"",MEDIAN(IF({lp}<>"",{lp})))', RATE)
        put(ws, f"C{r}", f'=IF(COUNT({lp})=0,"",_xlfn.MINIFS({lp},{lp},">0"))', RATE)
        put(ws, f"D{r}", job(f"C{r}")); ws.merge_cells(f"D{r}:E{r}")
        put(ws, f"F{r}", f'=IF(COUNT({lp})=0,"",_xlfn.MAXIFS({lp},{lp},">0"))', RATE)
        put(ws, f"G{r}", job(f"F{r}"))
        put(ws, f"H{r}", f"=COUNT({lp})", INT)
        put(ws, f"I{r}", f"=SUM({sf})", INT)
        arr(ws, f"J{r}", f'=IF(COUNT({spm})=0,"",MEDIAN(IF({spm}<>"",{spm})))', "#,##0.0")
        for col in "ABCDEFGHIJ":
            ws[f"{col}{r}"].border = THIN
    rl = r0 + 2 + len(UNIT_TYPES)
    put(ws, f"A{rl}", f'="Estimates with wages above standard ({STD_WAGES}): "&COUNTIF(tblLabor[wages_vs_std],">0.01")'
        f'&" of "&COUNTA(tblLabor[project])&". Iterations count toward rates."', font=F_NOTE)
    chart = BarChart(); chart.type, chart.style, chart.title = "bar", 10, "Median labor $/SF"
    chart.add_data(Reference(ws, min_col=2, min_row=r0 + 1, max_row=r0 + 1 + len(UNIT_TYPES)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=r0 + 2, max_row=r0 + 1 + len(UNIT_TYPES)))
    chart.legend, chart.height, chart.width = None, 6.5, 12
    ws.add_chart(chart, "M7")

    def idx_row(r, tbl, key, n, label_font=F_BASE):
        put(ws, f"A{r}", f'=IFERROR(IF(INDEX({tbl}[{key}],{n})="","",INDEX({tbl}[{key}],{n})),"")', font=label_font)
        for j, (h, f) in enumerate(zip(STAT_HEAD, STAT_FMT)):
            col = L(2 + j)
            put(ws, f"{col}{r}", f'=IF($A{r}="","",INDEX({tbl}[{h}],{n}))', f)
        for col in "ABCDEFGHIJ":
            ws[f"{col}{r}"].border = THIN

    h0 = rl + 3
    put(ws, f"A{h0}", '="Hit rate by project — "&C4&" (projects won ÷ won + lost; a project bid to several GCs counts once)"', font=F_SECTION)
    head(ws, h0 + 1, 1, [""] + ["Projects"] + STAT_HEAD[1:])
    r = h0 + 2
    cnt = lambda s: f'COUNTIFS(tblBidLog[project_first],1,tblBidLog[project_status],"{PSTAT[s]}",tblBidLog[in_period],1)'
    sm = lambda s: f'SUMIFS(tblBidLog[project_amount],tblBidLog[project_first],1,tblBidLog[project_status],"{PSTAT[s]}",tblBidLog[in_period],1)'
    put(ws, f"A{r}", "All projects", font=F_BOLD)
    for col, (v, f) in zip("BCDEFGHIJ", [(f"={cnt('submitted')}+{cnt('awarded')}+{cnt('lost')}+{cnt('closed')}", INT),
                                         (f"={cnt('awarded')}", INT), (f"={cnt('lost')}", INT),
                                         (f'=IF(C{r}+D{r}=0,"",C{r}/(C{r}+D{r}))', PCT), (f"={sm('awarded')}", MONEY),
                                         (f"={sm('lost')}", MONEY), (f'=IF(F{r}+G{r}=0,"",F{r}/(F{r}+G{r}))', PCT),
                                         (f"={cnt('submitted')}", INT), (f"={cnt('closed')}", INT)]):
        put(ws, f"{col}{r}", v, f); ws[f"{col}{r}"].border = THIN
    put(ws, f"A{r + 1}", "By job type", font=F_BOLD)
    for n in range(1, N_JOBTYPES + 1):
        idx_row(r + 1 + n, "tblJobTypes", "Job type", n)
    hr = r + 2 + N_JOBTYPES
    put(ws, f"A{hr}", '="Projects = open + won + lost + closed. Open (a GC still out) and closed (no result) are not in the rate. "'
        '&"$ uses each project\'s price once ("&COUNTIFS(tblBidLog[project_first],1,tblBidLog[project_status],"awarded",tblBidLog[project_amount],"")'
        '+COUNTIFS(tblBidLog[project_first],1,tblBidLog[project_status],"lost",tblBidLog[project_amount],"")'
        '&" won/lost projects have no amount). By GC counts every bid to that GC."', font=F_NOTE)

    g0 = hr + 2
    put(ws, f"A{g0}", '="Top GCs by bids (all-time order) — "&C4&" — each bid to the GC counts"', font=F_SECTION)
    head(ws, g0 + 1, 1, ["GC"] + STAT_HEAD)
    for n in range(1, N_TOPGC + 1):
        idx_row(g0 + 1 + n, "tblGCs", "GC", n)

    m0 = g0 + 3 + N_TOPGC
    put(ws, f"A{m0}", "Price movers — same supplier, same product, ≥ 3% vs a quote 30+ days earlier, latest quote in the last 90 days",
        font=F_SECTION)
    head(ws, m0 + 1, 1, ["Item", "Change", "Now", "Before", "Product", "Unit", "Supplier", "Quoted", "Prior", "Project"])
    for n in range(1, N_MOVERS + 1):
        r = m0 + 1 + n
        for col, key, f in zip("ABCDEFGHIJ", ["item", "change_pct", "unit_price", "prior_price", "product", "unit", "supplier",
                                              "quote_date", "prior_date", "project"],
                               [None, PCT, RATE, RATE, None, None, None, DATE, DATE, None]):
            put(ws, f"{col}{r}", f'=IFERROR(IF(INDEX(tblMovers[{key}],{n})="","",INDEX(tblMovers[{key}],{n})),"")', f)
            ws[f"{col}{r}"].border = THIN
    put(ws, f"A{m0 + 2 + N_MOVERS}", '=COUNTA(tblMovers[item])&" movers. Every latest price: _Knowledge/Price Book/Price Book Latest.csv. '
        'Refreshed with the Friday sweep."', font=F_NOTE)

    j0 = m0 + 4 + N_MOVERS
    put(ws, f"A{j0}", "Actual vs estimate — closed-out jobs (Job Cost tab)", font=F_SECTION)
    head(ws, j0 + 1, 1, ["", "Jobs", "Median man-day var", "Median labor $ var", "", "", "", "", "", ""])
    sc, st, mdv, lv = "tblJobCost[scope]", "tblJobCost[status]", "tblJobCost[mandays_var_pct]", "tblJobCost[labor_var_pct]"
    put(ws, f"A{j0 + 2}", "Final, job level", font=F_BOLD)
    put(ws, f"B{j0 + 2}", f'=COUNTIFS({sc},"job",{st},"final")', INT)
    arr(ws, f"C{j0 + 2}", f'=IF(COUNTIFS({sc},"job",{st},"final",{mdv},"<>")=0,"",MEDIAN(IF(({sc}="job")*({st}="final")*({mdv}<>""),{mdv})))', PCT)
    arr(ws, f"D{j0 + 2}", f'=IF(COUNTIFS({sc},"job",{st},"final",{lv},"<>")=0,"",MEDIAN(IF(({sc}="job")*({st}="final")*({lv}<>""),{lv})))', PCT)
    put(ws, f"A{j0 + 3}", "Positive = actual over estimate. Labor $ only where the accounting export is base wage; unit-type rows only "
        "for single-unit-type jobs or a manual split.", font=F_NOTE)
    ws.freeze_panes = "A6"

    for s_ in (dash, sh["By GC"], sh["By Job Type"]):
        s_.page_setup.orientation = "landscape"
        s_.sheet_properties.pageSetUpPr.fitToPage = True
        s_.page_setup.fitToWidth, s_.page_setup.fitToHeight = 1, 0
    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    wb.save(a.out)


PARTS, PART_MAX = 4, 18_000


def _dump(obj):
    return json.dumps(obj, ensure_ascii=False, separators=(",", ":"))


def part_paths(path):
    stem = path[:-4] if path.lower().endswith(".txt") else path
    return [f"{stem} {k}.txt" for k in range(1, PARTS + 1)]


def write_data(path, P):
    """Split the table rows into PARTS files of <= PART_MAX bytes. Returns (paths, problems)."""
    import hashlib
    tables = {n: {"columns": c, "rows": r} for n, (c, r) in P["T"].items()}
    set_id = hashlib.sha1(_dump(tables).encode()).hexdigest()[:12]
    parts = [{"part": k, "parts": PARTS, "set": set_id, "tables": {}} for k in range(1, PARTS + 1)]
    parts[-1].update(stamp=P["stamp"], about=P["about"])  # last part: usually tiny, so quiet nights upload ~2 KB
    k, problems = 0, []
    size = lambda d: len(_dump(d).encode()) + 1
    for name, t in tables.items():
        parts[k]["tables"].setdefault(name, {"columns": t["columns"], "rows": []})
        for row in t["rows"]:
            parts[k]["tables"][name]["rows"].append(row)
            if size(parts[k]) > PART_MAX:
                parts[k]["tables"][name]["rows"].pop()
                if not parts[k]["tables"][name]["rows"]:
                    del parts[k]["tables"][name]
                k += 1
                if k >= PARTS:
                    problems.append(f"data needs more than {PARTS} parts of {PART_MAX // 1000} KB — add a part (script + Flow 3)")
                    return [], problems
                parts[k]["tables"][name] = {"columns": t["columns"], "rows": [row]}
    paths = part_paths(path)
    for d, pth in zip(parts, paths):
        with open(pth, "w", encoding="utf-8") as f:
            f.write(_dump(d) + "\n")
    return paths, problems


# ---------- recalc + checks ----------
MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub RecalculateAndSave()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>"""


def recalc(path, timeout=120):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return "LibreOffice (soffice) not found — formulas were not recalculated"
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
    with tempfile.TemporaryDirectory(prefix="lo-profile-") as prof:
        url = "file://" + prof
        subprocess.run([soffice, "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                       capture_output=True, timeout=timeout, env=env)
        mdir = os.path.join(prof, "user", "basic", "Standard")
        if not os.path.isdir(mdir):
            return "LibreOffice profile not created — formulas were not recalculated"
        open(os.path.join(mdir, "Module1.xba"), "w").write(MACRO)
        before = os.stat(path).st_mtime_ns
        p = subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation={url}",
                            "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
                            os.path.abspath(path)], capture_output=True, text=True, timeout=timeout, env=env)
        if p.returncode != 0 or os.stat(path).st_mtime_ns == before:
            return f"LibreOffice did not rewrite the file (exit {p.returncode}) {p.stderr.strip()[:200]}"
    return None


ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A", "Err:")


def check(path, P):
    wf, wb = load_workbook(path), load_workbook(path, data_only=True)
    problems, nformula = [], 0
    for ws in wf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, ArrayFormula) or (isinstance(v, str) and v.startswith("=")):
                    nformula += 1
                    cv = wb[ws.title][c.coordinate].value
                    if isinstance(cv, str) and any(cv.startswith(e) for e in ERRORS):
                        problems.append(f"{ws.title}!{c.coordinate} = {cv}")
    d = wb["Dashboard"]
    row = next(r for r in range(1, 80) if d[f"A{r}"].value == "All projects")
    won = sum(v == "awarded" for v in P["proj"].values()); lost = sum(v == "lost" for v in P["proj"].values())
    if (d[f"C{row}"].value, d[f"D{row}"].value) != (won, lost):
        problems.append(f"hit rate cross-check: sheet {(d[f'C{row}'].value, d[f'D{row}'].value)} vs python {(won, lost)}")
    for u, name in UNIT_TYPES:
        vals = [num(r.get(f"{u}_labor_per_sf")) for r in P["bm"] if num(r.get(f"{u}_sf"))]
        vals = [v for v in vals if v is not None]
        rr = next(r for r in range(1, 40) if d[f"A{r}"].value == name)
        got, exp = d[f"B{rr}"].value, (round(statistics.median(vals), 3) if vals else None)
        if (got in ("", None)) != (exp is None) or (exp is not None and abs(float(got) - exp) > 0.005):
            problems.append(f"labor $/SF cross-check {name}: sheet {got} vs python {exp}")
    return nformula, problems


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--bid-log", required=True); ap.add_argument("--bid-metrics", required=True)
    ap.add_argument("--prices", default=""); ap.add_argument("--job-cost", default="")
    ap.add_argument("--out", default="Dashboard.xlsx"); ap.add_argument("--data", default="Dashboard Data.txt")
    ap.add_argument("--previous", default="", help="folder holding the current Dashboard Data parts (skip unchanged)")
    ap.add_argument("--period", default="YTD", choices=PERIODS)
    ap.add_argument("--stamp", default=dt.datetime.now().strftime("%Y-%m-%d %H:%M"))
    ap.add_argument("--no-recalc", action="store_true")
    a = ap.parse_args()
    P = prepare(a)
    build(a, P)
    paths, problems = write_data(a.data, P)
    nf = 0
    if not a.no_recalc:
        # verify on a copy with period = All (checks compare all-time counts)
        out, period = a.out, a.period
        a.out, a.period = out + ".check.xlsx", "All"
        build(a, P)
        err = recalc(a.out)
        if err:
            problems.append(err)
        else:
            nf, cp = check(a.out, P)
            problems += cp
        os.remove(a.out)
        a.out, a.period = out, period
    psz = ", ".join(f"{os.path.getsize(p) // 1024 + 1} KB" for p in paths)
    print(f"Dashboard: {a.out} ({os.path.getsize(a.out) // 1024} KB) + {len(paths)} data parts ({psz}) · "
          f"{len(P['bl'])} bids · {len(P['bm'])} estimates · {len(P['T']['tblGCs'][1])} GCs · {len(P['T']['tblMovers'][1])} movers · "
          f"{nf} formulas checked")
    for w in P["warn"]:
        print("⚠ " + w)
    for p in problems:
        print("✗ " + p)
    if problems:
        print("Upload: NO — fix the ✗ lines")
        return 1
    up = []
    for k, pth in enumerate(paths, 1):
        prev = os.path.join(a.previous, os.path.basename(pth)) if a.previous else None
        same = prev and os.path.exists(prev) and open(prev, encoding="utf-8").read().strip() == open(pth, encoding="utf-8").read().strip()
        if k == len(paths) or not same:
            up.append(pth)
    print("Upload: yes — text files, replace each (pass expectedBytes):")
    for pth in up:
        print(f"  {os.path.basename(pth)}  ({os.path.getsize(pth)} bytes)")
    if len(up) < len(paths):
        print(f"  ({len(paths) - len(up)} unchanged part(s) not uploaded)")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
